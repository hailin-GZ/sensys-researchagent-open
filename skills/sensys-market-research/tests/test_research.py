"""Behavioral integrity checks; fixed fixtures are not real market observations."""
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import research as r

HTML=b'''<html><head><title>Fixture title</title><meta name="author" content="Fixture Author"></head>
<body><div id="copyright_info">Original publisher <a href="https://example.org/original">Original</a></div>
<article><p>A fixture customer has a recording task and wants to retrieve a past decision accurately.</p>
<p>This fixture contains a second complete paragraph to test body extraction without inventing evidence.</p>
<img src="https://example.org/image.png"><img src="https://example.org/missing.png"></article></body></html>'''
from io import BytesIO
from PIL import Image
image_buffer=BytesIO()
Image.new('RGB',(2,2),'white').save(image_buffer,format='PNG')
PNG=image_buffer.getvalue()


class Integrity(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.root=Path(self.temp.name)/'task'
        r.init_task(self.root,{'question':'Analyze a fixed test fixture','decision_to_support':'Check integrity'},['synthesis'])
    def tearDown(self): self.temp.cleanup()
    def fetch(self,url,*args,**kwargs):
        if 'missing' in url: return b'missing',url,404,'text/plain'
        if url.endswith('.png'): return PNG,url,200,'image/png'
        return HTML,url,200,'text/html'
    def source(self):
        with patch.object(r,'download',side_effect=self.fetch):
            return r.capture(self.root,url='https://example.org/article',category='fixture')
    def result(self,source):
        block=r.read_json(self.root/source['artifacts']['blocks'])[0]
        quote='A fixture customer'
        e=r.add_evidence(self.root,source['source_id'],block['block_id'],quote)
        result=r.read_json(self.root/'result.json')
        result.update(summary='A fixed-fixture test.',support_reviewed=True,findings=[{'finding_id':'f1','kind':'fact','statement':'The fixture mentions a customer.','evidence_refs':[e['evidence_id']],'limitations':['Fixed test fixture; no market implication.']}])
        r.write_json(self.root/'result.json',result)
        return result
    def test_archive_raw_author_attribution_and_failed_image(self):
        s=self.source()
        self.assertEqual((self.root/s['artifacts']['raw']).read_bytes(),HTML)
        self.assertEqual(s['author'],'Fixture Author')
        self.assertTrue(s['attribution'][0]['links'])
        self.assertEqual([a['status'] for a in s['assets']],['saved','failed'])
        self.assertIn('https://example.org/missing.png',(self.root/s['artifacts']['markdown']).read_text())
    def test_complete_handoff_with_exact_quote(self):
        s=self.source();self.result(s)
        out=r.finalize(self.root)
        self.assertEqual(out['status'],'partial')
        self.assertTrue(out['validation']['ok'])
        self.assertIn('e0001',(self.root/'report.md').read_text())
    def test_quote_fabrication_rejected(self):
        s=self.source()
        with self.assertRaises(ValueError): r.add_evidence(self.root,s['source_id'],'p0001','Fabricated interview')
    def test_raw_tampering_rejected(self):
        s=self.source();self.result(s)
        (self.root/s['artifacts']['raw']).write_bytes(b'tampered')
        self.assertFalse(r.validate(self.root)['ok'])
    def test_markdown_tampering_rejected(self):
        s=self.source();self.result(s)
        (self.root/s['artifacts']['markdown']).write_text('tampered')
        self.assertFalse(r.validate(self.root)['ok'])
    def test_source_record_deletion_rejected(self):
        s=self.source();self.result(s)
        (self.root/s['source_record_path']).unlink()
        self.assertFalse(r.validate(self.root)['ok'])
    def test_asset_registry_deletion_rejected(self):
        s=self.source();self.result(s)
        (self.root/s['artifacts']['assets']).unlink()
        self.assertFalse(r.validate(self.root)['ok'])
    def test_source_manifest_metadata_tampering_rejected(self):
        s=self.source();self.result(s)
        records=r.read_json(self.root/'sources.json');records[0]['title']='Changed title';r.write_json(self.root/'sources.json',records)
        self.assertFalse(r.validate(self.root)['ok'])
    def test_truncated_image_not_saved(self):
        def broken(url,*args,**kwargs):
            if url.endswith('.png'): return PNG[:16],url,200,'image/png'
            return self.fetch(url,*args,**kwargs)
        with patch.object(r,'download',side_effect=broken): s=r.capture(self.root,url='https://example.org/article')
        self.assertEqual(s['assets'][0]['status'],'failed')
    def test_missing_markdown_image_rejected_even_with_fresh_hash(self):
        s=self.source();self.result(s)
        p=self.root/s['artifacts']['markdown'];p.write_text(p.read_text()+ '\n![chart](images/absent.png)',encoding='utf-8')
        records=r.read_json(self.root/'sources.json');records[0]['markdown_sha256']=r.sha(p.read_bytes())
        snapshot={k:v for k,v in records[0].items() if k!='source_record_sha256'}
        r.write_json(self.root/s['source_record_path'],snapshot)
        records[0]['source_record_sha256']=r.sha((self.root/s['source_record_path']).read_bytes());r.write_json(self.root/'sources.json',records)
        self.assertFalse(r.validate(self.root)['ok'])
    def test_empty_author_remains_unknown(self):
        data=HTML.replace(b'content="Fixture Author"',b'content=""')
        with patch.object(r,'download',return_value=(data,'https://example.org/a',200,'text/html')): s=r.capture(self.root,url='https://example.org/a',max_images=0)
        self.assertIsNone(s['author'])
    def test_failed_sources_in_report(self):
        s=self.source();self.result(s)
        with patch.object(r,'download',return_value=(b'Blocked','https://example.org/blocked',403,'text/plain')): failed=r.capture(self.root,url='https://example.org/blocked')
        r.finalize(self.root)
        self.assertIn(failed['source_id'],(self.root/'report.md').read_text())
    def test_unresolved_evidence_rejected(self):
        s=self.source();out=self.result(s);out['findings'][0]['evidence_refs']=['missing'];r.write_json(self.root/'result.json',out)
        self.assertFalse(r.validate(self.root)['ok'])
    def test_semantic_review_required(self):
        s=self.source();out=self.result(s);out['support_reviewed']=False;r.write_json(self.root/'result.json',out)
        self.assertFalse(r.validate(self.root)['ok'])
    def test_missing_selector_has_no_evidence(self):
        with patch.object(r,'download',side_effect=self.fetch): s=r.capture(self.root,url='https://example.org/a',selector='#absent')
        self.assertEqual(s['status'],'shell_only')
        with self.assertRaises(ValueError): r.add_evidence(self.root,s['source_id'],'p0001','customer')
    def test_no_body_challenge_not_success(self):
        with patch.object(r,'download',return_value=(b'<html><body>Verify you are human CAPTCHA</body></html>','https://example.org/a',200,'text/html')):
            s=r.capture(self.root,url='https://example.org/a')
        self.assertEqual(s['status'],'captcha')
    def test_duplicate_capture_reuses_source(self):
        s=self.source();again=self.source()
        self.assertEqual(s['source_id'],again['source_id'])
        self.assertTrue(again['reused'])
        self.assertEqual(len(r.read_json(self.root/'sources.json')),1)
    def test_content_duplicates_share_origin(self):
        s=self.source()
        with patch.object(r,'download',side_effect=self.fetch): second=r.capture(self.root,url='https://example.org/repost')
        self.assertEqual(second['duplicate_of'],s['source_id'])
        self.assertEqual(second['origin_id'],s['origin_id'])
    def test_path_escape_rejected(self):
        with self.assertRaises(ValueError): r.inside(self.root,'../outside')
    def test_private_url_rejected(self):
        with self.assertRaises(ValueError): r.public_url('http://127.0.0.1/private')
    def test_excerpt_requires_limitations(self):
        file=Path(self.temp.name)/'excerpt.md';file.write_text('A fixture customer has a task.',encoding='utf-8')
        s=r.capture(self.root,file=str(file),source_url='https://example.org/source',scope='excerpt',method='host_reader')
        out=self.result(s);out['findings'][0]['limitations']=[];r.write_json(self.root/'result.json',out)
        self.assertEqual(s['status'],'snippet_only');self.assertFalse(r.validate(self.root)['ok'])
    def test_office_and_csv_locators(self):
        from docx import Document
        from openpyxl import Workbook
        doc=Document();doc.add_paragraph('A supplied fixed test paragraph.');file=Path(self.temp.name)/'fixture.docx';doc.save(file)
        s=r.capture(self.root,file=str(file));self.assertEqual(r.read_json(self.root/s['artifacts']['blocks'])[0]['block_id'],'paragraph-0001')
        book=Workbook();book.active.append(['question','answer']);book.active.append(['fixture','=1+1']);file=Path(self.temp.name)/'fixture.xlsx';book.save(file)
        s=r.capture(self.root,file=str(file));blocks=r.read_json(self.root/s['artifacts']['blocks']);self.assertEqual(blocks[1]['row'],2);self.assertIn('=1+1',blocks[1]['text'])
    def test_completed_task_cannot_hide_decisive_gap(self):
        s=self.source();out=self.result(s);out.update(status='complete',missing_evidence=['Purchase behavior is unknown']);r.write_json(self.root/'result.json',out)
        self.assertFalse(r.validate(self.root)['ok'])


if __name__=='__main__': unittest.main(verbosity=2)
