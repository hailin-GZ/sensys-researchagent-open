"""Portable acquisition and evidence helpers; reasoning belongs to the host agent."""
from __future__ import annotations

import argparse
import csv
import hashlib
import ipaddress
import json
import os
import re
import shutil
import socket
import subprocess
import sys
import time
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urljoin, urlsplit, urlunsplit
import research_depth as depth

MODES = {'needs', 'synthesis', 'competitors', 'voices', 'demand', 'crowdfunding', 'sizing', 'trends', 'validation'}
KINDS = {'fact', 'source_opinion', 'inference', 'hypothesis', 'unknown'}
BAD_SOURCE = {'captcha', 'login_required', 'not_found', 'network_error', 'parse_error', 'shell_only', 'image_only'}
SKILL = Path(__file__).resolve().parents[1]
PARSER_VERSION = 6


def now():
    return datetime.now(timezone.utc).isoformat()


def sha(data):
    return hashlib.sha256(data).hexdigest()


def write_json(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix(path.suffix + '.tmp')
    temp.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    temp.replace(path)


def read_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8-sig'))


def cache_event(root, operation, status, **details):
    """Task-local, single-writer metrics; never imply provider prompt-cache hits."""
    path = Path(root) / 'cache-events.json'
    events = read_json(path) if path.exists() else []
    events.append(dict(operation=operation, status=status, checked_at=now(), **details))
    write_json(path, events[-1000:])


def age_seconds(timestamp):
    try:
        stamp = datetime.fromisoformat(timestamp)
        age = (datetime.now(timezone.utc) - stamp).total_seconds()
        return age if age >= 0 else None
    except (TypeError, ValueError):
        return None


def source_integrity(root, source):
    """Verify the reusable archive without requiring a finished result.json."""
    try:
        for field, key in [('raw','raw_sha256'),('markdown','markdown_sha256'),('blocks','blocks_sha256'),('assets','assets_sha256'),('source_record','source_record_sha256')]:
            rel = source.get('source_record_path') if field == 'source_record' else source['artifacts'][field]
            file = inside(root, rel)
            if not file.is_file() or sha(file.read_bytes()) != source[key]: return False
        record = read_json(inside(root, source['source_record_path']))
        if record != {k:v for k,v in source.items() if k != 'source_record_sha256'}: return False
        if read_json(inside(root,source['artifacts']['assets'])) != source.get('assets',[]): return False
        for asset in source.get('assets',[]):
            if asset['status'] == 'saved':
                file = inside(root, (Path(source['artifacts']['markdown']).parent / asset['path']).as_posix())
                if not file.is_file() or sha(file.read_bytes()) != asset['sha256'] or not image_type(file.read_bytes()): return False
        return True
    except (KeyError, ValueError, TypeError, OSError):
        return False


def inside(root, value):
    candidate = (Path(root) / value).resolve()
    if not candidate.is_relative_to(Path(root).resolve()):
        raise ValueError('Artifact path escapes task root')
    return candidate


def public_url(url):
    parts = urlsplit(url)
    if parts.scheme not in {'http', 'https'} or not parts.hostname or parts.username or parts.password:
        raise ValueError('A public HTTP(S) URL without credentials is required')
    try:
        literal = ipaddress.ip_address(parts.hostname)
    except ValueError:
        literal = None
    if literal and not literal.is_global:
        raise ValueError('Non-public literal IP destinations are not supported')
    for address in socket.getaddrinfo(parts.hostname, parts.port or (443 if parts.scheme == 'https' else 80)):
        resolved = ipaddress.ip_address(address[4][0])
        fake_dns = os.environ.get('SENSYS_ALLOW_BENCHMARK_DNS') == '1' and not literal and resolved in ipaddress.ip_network('198.18.0.0/15')
        if not resolved.is_global and not fake_dns:
            raise ValueError('Non-public network destinations are not supported')
    return url


def normalized_url(url):
    p = urlsplit(url)
    return urlunsplit((p.scheme.lower(), p.netloc.lower(), p.path, p.query, ''))


def download(url, timeout=15, max_bytes=25_000_000):
    import requests
    session = requests.Session()
    session.trust_env = False
    session.headers['User-Agent'] = 'Mozilla/5.0 (compatible; SENSYSResearch/1.0)'
    start = time.monotonic()
    for _ in range(6):
        public_url(url)
        with session.get(url, timeout=timeout, allow_redirects=False, stream=True) as response:
            if response.is_redirect:
                url = urljoin(url, response.headers['Location'])
                continue
            chunks, size = [], 0
            for chunk in response.iter_content(65536):
                size += len(chunk)
                if size > max_bytes or time.monotonic() - start > timeout * 2:
                    raise ValueError('Download byte/time budget exceeded')
                chunks.append(chunk)
            return b''.join(chunks), url, response.status_code, response.headers.get('Content-Type', '')
    raise ValueError('Redirect limit exceeded')


def init_task(root, brief, modes, research_depth=None):
    root = Path(root)
    if (root / 'brief.json').exists():
        raise ValueError('Task already exists; reuse it or choose another directory')
    selected = list(dict.fromkeys(modes))
    if not selected or set(selected) - MODES:
        raise ValueError('Select valid research modes')
    chosen_depth = research_depth or brief.get('research_depth', 'standard')
    if research_depth and brief.get('research_depth') and research_depth != brief['research_depth']:
        raise ValueError('CLI depth conflicts with brief research_depth; select one explicitly')
    depth.effective_limits(chosen_depth, brief)
    root.mkdir(parents=True, exist_ok=True)
    brief = dict(brief)
    brief.update(task_id=brief.get('task_id') or f'research-{datetime.now(timezone.utc):%Y%m%dT%H%M%S}', selected_modes=selected, research_depth=chosen_depth, created_at=now())
    write_json(root / 'brief.json', brief)
    write_json(root / 'sources.json', [])
    write_json(root / 'evidence.json', [])
    write_json(root / 'discoveries.json', [])
    write_json(root / 'result.json', {'task_id':brief['task_id'], 'status':'partial', 'executed_scope':selected, 'research_depth':chosen_depth,
        'summary':'', 'findings':[], 'conflicts':[], 'missing_evidence':[], 'failed_sources':[],
        'assumptions':[], 'next_steps':[], 'support_reviewed':False,
        'source_manifest_path':'sources.json', 'evidence_manifest_path':'evidence.json', 'report_path':'report.md'})
    depth.create(root, brief, chosen_depth)
    return brief


@depth.serial_task
def discover(root, query, limit=8, backend='auto', max_age=3600, refresh=False):
    if type(limit) is not int or limit < 1:
        raise ValueError('Search result limit must be a positive integer')
    limit = depth.capped(root, 'results_per_query', limit)
    records = read_json(Path(root) / 'discoveries.json')
    if max_age < 0: raise ValueError('Cache max age must be non-negative')
    key = sha(json.dumps([query, backend, limit], ensure_ascii=False).encode('utf-8'))
    for previous in reversed(records):
        age = age_seconds(previous.get('retrieved_at'))
        digest = sha(json.dumps(previous.get('results',[]), sort_keys=True, ensure_ascii=False).encode('utf-8'))
        if not refresh and max_age > 0 and previous.get('cache_key') == key and previous.get('status') == 'leads_found' and age is not None and age <= max_age and previous.get('results_sha256') == digest:
            cache_event(root,'discover','hit',cache_key=key,age_seconds=round(age,3))
            return {**previous,'reused':True,'cache_age_seconds':age}
    depth.reserve(root, 'search_queries')
    cache_event(root,'discover','miss',cache_key=key,forced_refresh=refresh)
    item = {'query':query, 'provider':'ddgs', 'backend':backend, 'retrieved_at':now(), 'results':[], 'effective_limit':limit}
    try:
        child = subprocess.run([sys.executable,str(SKILL/'scripts/search_worker.py'),'--query',query,'--limit',str(limit),'--backend',backend],capture_output=True,text=True,encoding='utf-8',timeout=40)
        diagnostic=getattr(child,'stderr','') or ''
        try:item['provider_attempts']=json.loads(diagnostic.strip().splitlines()[-1])
        except (ValueError,IndexError):item['provider_attempts']=[{'diagnostic':diagnostic[-800:]}] if diagnostic else []
        if child.returncode:
            raise RuntimeError('Search provider failed; see provider_attempts')
        for result in json.loads(child.stdout):
            item['results'].append({'url':result.get('href'), 'title':result.get('title'), 'snippet':result.get('body')})
        item['status'] = 'leads_found' if item['results'] else 'no_results'
    except Exception as exc:
        item.update(status='failed', error=type(exc).__name__, diagnostic=str(exc)[-500:])
    item['cache_key'] = key
    item['results_sha256'] = sha(json.dumps(item['results'],sort_keys=True,ensure_ascii=False).encode('utf-8'))
    records.append(item)
    write_json(Path(root) / 'discoveries.json', records)
    return item


def image_type(data):
    from io import BytesIO
    from xml.etree import ElementTree
    from PIL import Image
    try:
        if b'<svg' in data[:1000].lower():
            if b'<!DOCTYPE' in data.upper() or b'<!ENTITY' in data.upper(): return None
            node = ElementTree.fromstring(data)
            return 'svg' if node.tag.split('}')[-1] == 'svg' else None
        with Image.open(BytesIO(data)) as decoded:
            extension = {'PNG':'png','JPEG':'jpg','GIF':'gif','WEBP':'webp'}.get(decoded.format)
            decoded.verify()
        return extension
    except Exception:
        return None


def html_document(data, url, folder, selector=None, max_images=30, timeout=15,stop_requested=None):
    from bs4 import BeautifulSoup
    from markdownify import markdownify
    soup = BeautifulSoup(data, 'html.parser')
    title = soup.title.get_text(' ', strip=True) if soup.title else ''
    meta = lambda key: next((tag.get('content') for tag in soup.find_all('meta') if tag.get('name') == key or tag.get('property') == key), None)
    title = meta('og:title') or title
    author = meta('author') or None
    publisher = meta('og:site_name') or None
    date = meta('article:published_time') or meta('date') or None
    jats=soup.find('article-meta') is not None and soup.find('article-title') is not None
    if jats:
        title=soup.find('article-title').get_text(' ',strip=True)
        journal=soup.find('journal-title');publisher=journal.get_text(' ',strip=True) if journal else publisher
        names=soup.select('article-meta contrib[contrib-type="author"] name')
        author='; '.join(n.get_text(' ',strip=True) for n in names) or author
    for tag in soup.find_all('script',type='application/ld+json'):
        try:
            pending = [json.loads(tag.string or tag.get_text())]
            while pending:
                item = pending.pop()
                if isinstance(item,list): pending.extend(item); continue
                if not isinstance(item,dict): continue
                pending.extend(value for value in item.values() if isinstance(value,(dict,list)))
                types = item.get('@type',[])
                types = [types] if isinstance(types,str) else types
                if isinstance(types,list) and any(t in {'Article','NewsArticle','BlogPosting','Review'} for t in types if isinstance(t,str)):
                    candidate = item.get('author')
                    if isinstance(candidate,list): candidate = candidate[0] if candidate else None
                    if isinstance(candidate,dict): author = author or candidate.get('name')
        except (ValueError,TypeError): pass
    full_text = soup.get_text(' ', strip=True)
    status, warning = 'partial_body', ['Body coverage requires review']
    blocked = full_text[:3000].lower()
    if urlsplit(url).hostname == 'mp.weixin.qq.com' and soup.select_one('#js_content') is None and (
        'Account moved' in title or 'This public account has moved' in full_text or 'This account was voluntarily closed' in full_text
    ):
        return {'title':title, 'status':'shell_only', 'blocks':[], 'assets':[], 'warnings':['Account migration/retirement notice; original article not obtained']}, ''
    if soup.select_one('#challenge-form, #challenge-running, .cf-error-details') is not None or title in {'Just a moment...', 'Attention Required! | Cloudflare'}:
        return {'title':title, 'status':'captcha', 'blocks':[], 'assets':[], 'warnings':['Security challenge; no body evidence']}, ''
    if len(full_text)<1000 and ('verify you are human' in blocked or 'complete the captcha' in blocked or 'Enter the verification code' in blocked or 'Security validation' in blocked):
        return {'title':title, 'status':'captcha', 'blocks':[], 'assets':[], 'warnings':['Challenge page; no body evidence']}, ''
    if len(full_text) < 500 and ('sign in to continue' in blocked or 'login required' in blocked):
        return {'title':title, 'status':'login_required', 'blocks':[], 'assets':[], 'warnings':['Login page; no body evidence']}, ''
    if selector:
        chosen = selector
        body = soup.select_one(selector)
    else:
        chosen = '#js_content'
        body = soup.select_one('#js_content')
        if body is not None and len(body.get_text(' ',strip=True))<80:body=None
        if body is None:
            for candidate in ('main','[role="main"]'):
                node=soup.select_one(candidate)
                if node is not None and len(node.get_text(' ',strip=True))>=80:
                    chosen=candidate;body=node;break
        if body is None:
            articles = soup.select('article')
            body = articles[0] if len(articles) == 1 else None
            chosen = 'article' if body is not None else 'body'
            if len(articles) > 1: warning.append('Multiple article cards; use a reviewed selector instead of treating the first card as the whole page')
    if body is not None:
        status = 'complete_body'
        warning = ['Selected body extracted; pagination, completeness and attribution require host review']
    elif selector:
        body = soup.new_tag('div')
        status, warning = 'shell_only', ['Requested body selector not found']
    else:
        body = soup.body or soup
    if 'mp.weixin.qq.com' == urlsplit(url).hostname:
        account = soup.select_one('#js_name')
        publisher = account.get_text(' ', strip=True) if account else publisher
        for pattern in [r'var\s+ct\s*=\s*["\x27]?(\d{10})', r'create_time\s*[=:]\s*["\x27]?(\d{10})']:
            match = re.search(pattern, data.decode('utf-8',errors='replace'))
            if match:
                date = datetime.fromtimestamp(int(match[1]), timezone.utc).isoformat()
                break
    attribution = []
    for node in soup.select('#copyright_info, .original_primary_card, #js_share_source'):
        attribution.append({'text':node.get_text(' ',strip=True), 'links':[urljoin(url,a['href']) for a in node.find_all('a',href=True)]})
    for node in body(['script','style','nav','footer','header','aside']):
        node.decompose()
    assets, seen = [], {}
    for img in body.find_all('img') + body.find_all('video',poster=True):
        target_attr = 'poster' if img.name == 'video' else 'src'
        original = img.get('poster') if img.name == 'video' else img.get('data-src') or img.get('data-original') or img.get('src')
        if not original:
            warning.append('Image tag without a retrievable source')
            continue
        original = urljoin(url, original)
        if original in seen:
            img[target_attr] = seen[original]
            continue
        asset = {'original_url':original, 'status':'omitted_budget', 'path':None, 'sha256':None, 'bytes':0}
        stopped=bool(stop_requested and stop_requested())
        if stopped:asset['status']='omitted_stop'
        if len(assets) < max_images and not stopped:
            try:
                raw, final, http, ctype = download(original, timeout=timeout, max_bytes=10_000_000)
                extension = image_type(raw)
                if http != 200 or not extension:
                    raise ValueError('Invalid image response')
                file = folder / 'images' / f'{sha(raw)[:20]}.{extension}'
                file.parent.mkdir(parents=True, exist_ok=True)
                file.write_bytes(raw)
                asset.update(status='saved', path=file.relative_to(folder).as_posix(), sha256=sha(raw), bytes=len(raw), final_url=final, media_type=extension)
            except Exception as exc:
                asset.update(status='failed', error=type(exc).__name__)
        target = asset['path'] or original
        seen[original] = target
        img[target_attr] = target
        for key in ['data-src','data-original','srcset']:
            img.attrs.pop(key,None)
        assets.append(asset)
    if urlsplit(url).hostname == 'mp.weixin.qq.com':
        from wechat_converter import HTMLToMarkdownParser
        class Registry:
            def register(self, source): return source
        parser = HTMLToMarkdownParser(Registry())
        parser.feed(str(body))
        markdown = parser.get_markdown()
    else:
        markdown = markdownify(str(body), heading_style='ATX')
    blocks = [{'block_id':f'p{i:04d}', 'text':p.strip()} for i,p in enumerate(re.split(r'\n\s*\n',markdown),1) if p.strip()]
    if not blocks or sum(len(b['text']) for b in blocks) < 80:
        status = 'shell_only'
    if jats and soup.find('graphic') is not None and status not in BAD_SOURCE:
        status='partial_body'
        warning.append('Original full-text XML includes figures not visually reviewed; extracted text does not verify image-only findings')
    return {'title':title, 'author':author, 'publisher':publisher, 'published_at':date, 'status':status,
        'blocks':blocks, 'assets':assets, 'warnings':warning, 'attribution':attribution, 'selector':chosen}, markdown


def parse_local(data, suffix):
    from io import BytesIO, StringIO
    blocks, warnings = [], []
    if suffix == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(data))
        for page, item in enumerate(reader.pages,1):
            text = item.extract_text() or ''
            blocks.append({'block_id':f'page-{page:04d}', 'text':text, 'pdf_page':page})
            if not text.strip(): warnings.append(f'Page {page} requires visual/OCR review')
        warnings.append('Original PDF preserves figures; extracted text does not verify charts')
    elif suffix == '.docx':
        from docx import Document
        doc = Document(BytesIO(data))
        blocks = [{'block_id':f'paragraph-{i:04d}','text':p.text} for i,p in enumerate(doc.paragraphs,1) if p.text.strip()]
        for ti, table in enumerate(doc.tables,1):
            for ri,row in enumerate(table.rows,1):
                blocks.append({'block_id':f'table-{ti}-row-{ri}','text':' | '.join(cell.text for cell in row.cells)})
        warnings.append('DOCX paragraph/table locations; original file preserves layout and images')
    elif suffix in {'.xlsx','.xlsm'}:
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        for sheet in book:
            for ri,row in enumerate(sheet.iter_rows(values_only=True),1):
                if any(v is not None for v in row):
                    blocks.append({'block_id':f'sheet-{book.sheetnames.index(sheet.title)+1}-row-{ri}', 'sheet':sheet.title, 'row':ri, 'text':json.dumps(list(row),ensure_ascii=False,default=str)})
        warnings.append('Formulas preserved as text; no recalculation or cached-result validation')
    elif suffix in {'.csv','.tsv'}:
        text = data.decode('utf-8-sig')
        for ri,row in enumerate(csv.reader(StringIO(text),delimiter='\t' if suffix == '.tsv' else ','),1):
            blocks.append({'block_id':f'row-{ri:04d}', 'text':json.dumps(row,ensure_ascii=False)})
    elif suffix == '.json':
        obj = json.loads(data.decode('utf-8-sig'))
        items = obj if isinstance(obj,list) else [obj]
        blocks = [{'block_id':f'record-{i:04d}', 'text':json.dumps(item,ensure_ascii=False,indent=2)} for i,item in enumerate(items,1)]
    elif suffix in {'.png','.jpg','.jpeg','.gif','.webp'}:
        return [], 'image_only', ['Original image preserved; host vision/OCR is required before citing its contents']
    elif suffix in {'.md','.txt'}:
        text = data.decode('utf-8-sig')
        blocks = [{'block_id':f'p{i:04d}','text':p.strip()} for i,p in enumerate(re.split(r'\n\s*\n',text),1) if p.strip()]
    else:
        raise ValueError('Unsupported file type; use an actual document/OCR tool and import its output')
    return blocks, 'structured_data' if suffix in {'.json','.csv','.tsv','.xlsx','.xlsm'} else 'supplied_text', warnings


@depth.serial_task
def capture(root, url=None, file=None, category='web', selector=None, max_images=30, timeout=15, origin_id=None, source_url=None, scope='full', method=None, max_age=86400, refresh=False, reserved=False,stop_requested=None):
    root = Path(root)
    sources = read_json(root / 'sources.json')
    if max_images < 0 or timeout < 1:
        raise ValueError('Image budget must be non-negative and timeout positive')
    max_images = depth.capped(root, 'images_per_source', max_images)
    if file and Path(file).is_file() and depth.load(root):
        if Path(file).stat().st_size > depth.profiles()['resource_policy']['local_import_max_bytes']:
            raise ValueError('Local import exceeds the 25 MB resource ceiling; use bounded external extraction and retain original provenance')
    if max_age < 0: raise ValueError('Cache max age must be non-negative')
    target = normalized_url(url or source_url) if url or source_url else str(Path(file).resolve())
    file_digest = sha(Path(file).read_bytes()) if file and Path(file).is_file() else None
    request = dict(target=target,selector=selector,max_images=max_images,scope=scope,
                   method=method or ('public_http' if url else 'supplied_file'),origin_id=origin_id,
                   file_sha256=file_digest,parser_version=PARSER_VERSION)
    key = sha(json.dumps(request,sort_keys=True,ensure_ascii=False).encode('utf-8'))
    for previous in reversed(sources):
        age = age_seconds(previous.get('retrieved_at'))
        if not refresh and max_age > 0 and previous.get('capture_key') == key and previous['status'] not in BAD_SOURCE and age is not None and age <= max_age and source_integrity(root,previous):
            cache_event(root,'capture','hit',source_id=previous['source_id'],age_seconds=round(age,3))
            return {**previous, 'reused':True,'cache_age_seconds':age}
    if not reserved:depth.reserve(root, 'captures')
    cache_event(root,'capture','miss',cache_key=key,forced_refresh=refresh)
    sid = f's{len(sources)+1:04d}'
    folder = root / 'sources' / sid
    folder.mkdir(parents=True, exist_ok=True)
    record = {'source_id':sid, 'requested_url':target, 'final_url':target, 'category':category,
        'retrieved_at':now(), 'acquisition_method':method or ('public_http' if url else 'supplied_file'),
        'origin_id':origin_id or target, 'duplicate_of':None, 'title':None, 'author':None,
        'publisher':None, 'published_at':None, 'status':'network_error', 'warnings':[], 'artifacts':{}, 'assets':[]}
    record['network_configuration'] = {'ambient_proxy_credentials':False,'benchmark_dns_compatibility':os.environ.get('SENSYS_ALLOW_BENCHMARK_DNS') == '1'}
    record['capture_key'] = key
    record['capture_request'] = request
    markdown, blocks = '', []
    try:
        if url:
            data, final, http, ctype = download(url, timeout)
            record.update(final_url=final,http_status=http,content_type=ctype)
            suffix = '.pdf' if data.startswith(b'%PDF-') else '.json' if 'json' in ctype else '.html'
        else:
            data = Path(file).read_bytes()
            suffix = Path(file).suffix.lower()
            record['title'] = Path(file).name
            http = 200
        raw_path = folder / f'raw{suffix}'
        raw_path.write_bytes(data)
        record['raw_sha256'] = sha(data)
        record['artifacts']['raw'] = raw_path.relative_to(root).as_posix()
        # Identical failed responses are not a shared real-world source origin.
        duplicate = next((s for s in sources if http==200 and s['status'] not in BAD_SOURCE and s.get('raw_sha256') == record['raw_sha256']),None)
        if duplicate:
            record.update(duplicate_of=duplicate['source_id'],origin_id=duplicate['origin_id'])
        if http != 200:
            record['status'] = 'not_found' if http == 404 else 'login_required' if http == 401 else 'network_error'
            record['warnings'] = [f'HTTP {http}; raw error response retained']
        elif suffix in {'.html','.htm'}:
            parsed, markdown = html_document(data,record['final_url'],folder,selector,max_images,timeout,stop_requested)
            blocks = parsed.pop('blocks')
            record.update(parsed)
        else:
            blocks, status, warnings = parse_local(data,suffix)
            if url and suffix == '.pdf':
                status = 'partial_body' if any('requires visual/OCR' in warning for warning in warnings) else 'complete_body'
            record.update(status=status,warnings=warnings)
            markdown = '\n\n'.join(b['text'] for b in blocks)
        if source_url:
            record['source_url'] = source_url
        if scope == 'excerpt' and record['status'] not in BAD_SOURCE:
            record['status'] = 'snippet_only'
            record['warnings'].append('Imported excerpt; not a complete original body')
    except Exception as exc:
        record.update(status='parse_error' if record['artifacts'].get('raw') else 'network_error', error=type(exc).__name__)
    md_path = folder / 'content.md'
    md_path.write_text(markdown,encoding='utf-8')
    write_json(folder / 'blocks.json', blocks)
    write_json(folder / 'assets.json',record['assets'])
    record['artifacts'].update(markdown=md_path.relative_to(root).as_posix(),blocks=(folder/'blocks.json').relative_to(root).as_posix(),assets=(folder/'assets.json').relative_to(root).as_posix())
    record['markdown_sha256'] = sha(md_path.read_bytes())
    record['blocks_sha256'] = sha((folder/'blocks.json').read_bytes())
    record['assets_sha256'] = sha((folder/'assets.json').read_bytes())
    record['source_record_path'] = (folder/'source.json').relative_to(root).as_posix()
    write_json(folder/'source.json',record)
    record['source_record_sha256'] = sha((folder/'source.json').read_bytes())
    sources.append(record)
    write_json(root/'sources.json',sources)
    return record


def add_evidence(root, source_id, block_id, quote, role='observation', context=None):
    root = Path(root)
    sources = {s['source_id']:s for s in read_json(root/'sources.json')}
    source = sources[source_id]
    if source['status'] in BAD_SOURCE:
        raise ValueError('Source has no usable text evidence')
    blocks = {b['block_id']:b for b in read_json(inside(root,source['artifacts']['blocks']))}
    if not quote.strip() or quote not in blocks[block_id]['text']:
        raise ValueError('Quote is not an exact substring of the captured block')
    evidence = read_json(root/'evidence.json')
    item = {'evidence_id':f'e{len(evidence)+1:04d}', 'source_id':source_id, 'block_id':block_id, 'quote':quote, 'evidence_role':role, 'context':context}
    evidence.append(item)
    write_json(root/'evidence.json',evidence)
    return item


def validate(root):
    root = Path(root)
    errors, warnings = [], []
    result = read_json(root/'result.json')
    sources_list = read_json(root/'sources.json')
    sources = {s['source_id']:s for s in sources_list}
    evidence_list = read_json(root/'evidence.json')
    evidence = {e['evidence_id']:e for e in evidence_list}
    if len(sources) != len(sources_list): errors.append('Duplicate source IDs')
    if len(evidence) != len(evidence_list): errors.append('Duplicate evidence IDs')
    blocks_cache = {}
    for sid,s in sources.items():
        for field,key in [('raw','raw_sha256'),('markdown','markdown_sha256'),('blocks','blocks_sha256'),('assets','assets_sha256'),('source_record','source_record_sha256')]:
            path = s.get('source_record_path') if field == 'source_record' else s['artifacts'].get(field)
            if not path:
                if field == 'raw' and s['status'] == 'network_error': continue
                errors.append(f'{sid}: missing {field} artifact')
                continue
            try:
                p = inside(root,path)
                if not p.is_file() or sha(p.read_bytes()) != s[key]: errors.append(f'{sid}: {field} missing or hash mismatch')
            except (ValueError,KeyError): errors.append(f'{sid}: invalid {field} path/hash')
        try:
            if read_json(inside(root,s['artifacts']['assets'])) != s.get('assets',[]): errors.append(f'{sid}: asset registry differs from source manifest')
            expected_record = {key:value for key,value in s.items() if key != 'source_record_sha256'}
            if read_json(inside(root,s['source_record_path'])) != expected_record: errors.append(f'{sid}: source record differs from manifest')
        except Exception: errors.append(f'{sid}: missing/invalid source or asset registry')
        try:
            blocks_cache[sid] = {b['block_id']:b for b in read_json(inside(root,s['artifacts']['blocks']))}
        except Exception: errors.append(f'{sid}: invalid blocks')
        try:
            md_file = inside(root,s['artifacts']['markdown'])
            for target in re.findall(r'!\[[^\]]*\]\(([^\s)]+)',md_file.read_text(encoding='utf-8')):
                target = target.strip('<>')
                if urlsplit(target).scheme in {'http','https','data'}: continue
                image_file = inside(root,(Path(s['artifacts']['markdown']).parent/target).as_posix())
                if not image_file.is_file() or not image_type(image_file.read_bytes()): errors.append(f'{sid}: unresolved local Markdown image')
        except Exception: errors.append(f'{sid}: invalid Markdown image reference')
        for asset in s.get('assets',[]):
            if asset['status'] == 'saved':
                try:
                    p = inside(root, str(Path(s['artifacts']['markdown']).parent / asset['path']))
                    if not p.is_file() or sha(p.read_bytes()) != asset['sha256'] or not image_type(p.read_bytes()): errors.append(f'{sid}: invalid saved image')
                except Exception: errors.append(f'{sid}: invalid asset path')
            else: warnings.append(f"{sid}: asset {asset['status']}")
        if s['status'] in BAD_SOURCE or s['status'] == 'partial_body': warnings.append(f"{sid}: {s['status']}")
    for eid,e in evidence.items():
        s = sources.get(e.get('source_id'))
        b = blocks_cache.get(e.get('source_id'),{}).get(e.get('block_id'))
        if not s or s['status'] in BAD_SOURCE or not b or not e.get('quote') or e['quote'] not in b['text']:
            errors.append(f'{eid}: missing/unusable source, block or exact quote')
    ids = set()
    for f in result.get('findings',[]):
        fid = f.get('finding_id')
        if not fid or fid in ids: errors.append('Missing/duplicate finding ID')
        ids.add(fid)
        if f.get('kind') not in KINDS or not f.get('statement'): errors.append(f'{fid}: invalid kind or statement')
        refs = f.get('evidence_refs',[])
        if f.get('kind') in {'fact','source_opinion','inference'} and not refs: errors.append(f'{fid}: evidence required')
        if f.get('kind') in {'hypothesis','unknown'} and not refs and not f.get('limitations'): errors.append(f'{fid}: explain unsupported hypothesis/unknown')
        for ref in refs:
            if ref not in evidence: errors.append(f'{fid}: unresolved evidence {ref}')
            elif evidence[ref].get('source_id') in sources and sources[evidence[ref]['source_id']]['status'] in {'partial_body','snippet_only'} and not f.get('limitations'):
                errors.append(f'{fid}: partial evidence requires limitations')
    if result.get('status') not in {'complete','partial','needs_input','failed'}: errors.append('Invalid task status')
    if result.get('task_id') != read_json(root/'brief.json')['task_id']: errors.append('Task ID mismatch')
    if result.get('findings') and result.get('support_reviewed') is not True: errors.append('Host semantic support review is required')
    if result.get('status') == 'complete' and result.get('missing_evidence'): errors.append('Unresolved decisive gaps require partial status')
    try:
        errors.extend(depth.check_review(root, result))
    except (ValueError, KeyError, TypeError, OSError) as exc:
        errors.append('Invalid research profile/budget: ' + str(exc))
    if (root/'research-plan.json').exists():
        from capabilities import check_plan
        errors.extend(check_plan(root,read_json(root/'research-plan.json'),complete=result.get('status')=='complete'))
    if not result.get('summary'): errors.append('Result summary is required')
    checked = {'checked_at':now(), 'ok':not errors, 'errors':errors, 'warnings':list(dict.fromkeys(warnings)), 'sources':len(sources), 'evidence':len(evidence), 'findings':len(result.get('findings',[])), 'semantic_check':'Declared host review; not an automated entailment check'}
    write_json(root/'validation.json',checked)
    return checked


def finalize(root):
    root = Path(root)
    checked = validate(root)
    if not checked['ok']: raise ValueError('Validation failed: ' + '; '.join(checked['errors']))
    result = read_json(root/'result.json')
    evidence = {e['evidence_id']:e for e in read_json(root/'evidence.json')}
    sources = {s['source_id']:s for s in read_json(root/'sources.json')}
    failures = list(result.get('failed_sources',[]))
    for sid,s in sources.items():
        if s['status'] in BAD_SOURCE and sid not in failures and not any(isinstance(item,dict) and item.get('source_id') == sid for item in failures): failures.append(sid)
    result['failed_sources'] = failures
    lines = ['# Research Result', '', f"Status: {result['status']}", '']
    budget = depth.status(root)
    if budget.get('research_depth'):
        lines.extend([f"Research depth: {budget['research_depth']}", '', 'Acquisition usage / ceiling: ' + ', '.join(f"{key} {budget['used'][key]}/{budget['limits'][key]}" for key in depth.COUNTERS), '', budget['host_limits'], ''])
        if budget['blocked_operations']:
            lines.extend(['Acquisition ceiling reached; see research-budget.json for blocked operations.', ''])
    lines.extend([result['summary'], '', '## Findings', ''])
    for f in result['findings']:
        refs = []
        for eid in f.get('evidence_refs',[]):
            e = evidence[eid]; s = sources[e['source_id']]
            target = s.get('source_url') or (s['final_url'] if s['acquisition_method'] == 'public_http' else s['artifacts']['raw'])
            refs.append(f"[{eid}: {e['block_id']}]({target})")
        lines.append(f"- **{f['finding_id']} / {f['kind']}**: {f['statement']} {' '.join(refs)}")
        if f.get('limitations'): lines.append('  Limitations: ' + '; '.join(f['limitations']))
    if (root/'research-plan.json').exists():
        plan = read_json(root/'research-plan.json')
        lines.extend(['','## Evidence coverage','', '| Role | Status | Sources | Remaining gap |','|---|---|---|---|'])
        for role in plan.get('roles',[]):
            cells = [role['role'],role['status'],', '.join(role.get('source_ids',[])),role.get('gap','')]
            lines.append('| '+' | '.join(str(c).replace('|','\\|').replace('\n',' ') for c in cells)+' |')
    if result.get('depth_review'):
        lines.extend(['', '## Depth review', ''])
        for check in result['depth_review']['checks']:
            lines.append(f"- {check['check']} / {check['status']}: {check['reason']} (findings: {', '.join(check.get('finding_refs', [])) or 'none'})")
    for key in ['assumptions','conflicts','missing_evidence','failed_sources','next_steps']:
        if result.get(key):
            lines.extend(['', '## ' + key.replace('_',' ').title(), ''])
            lines.extend('- ' + (item if isinstance(item,str) else json.dumps(item,ensure_ascii=False)) for item in result[key])
    lines.extend(['', '## Evidence index', ''])
    for eid,e in evidence.items():
        s=sources[e['source_id']]
        lines.extend([f"- {eid}: [{e['source_id']}]({s['artifacts']['markdown']}) / {e['block_id']}: {e['quote']}"])
    (root/'report.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')
    result['artifact_paths'] = ['brief.json','sources.json','evidence.json','discoveries.json','report.md','validation.json']
    result['artifact_paths'] += [name for name in ['research-plan.json','research-budget.json','cache-events.json','handoff.json'] if (root/name).is_file()]
    result['finalized_at']=now()
    result.pop('invalidated_artifacts', None)
    write_json(root/'result.json',result)
    return {'status':result['status'],'report_path':str(root/'report.md'),'result_path':str(root/'result.json'),'validation':checked}


def doctor():
    import importlib.metadata
    result = {}
    for package in ['requests','beautifulsoup4','markdownify','pypdf','ddgs','python-docx','openpyxl','Pillow']:
        try: result[package] = importlib.metadata.version(package)
        except importlib.metadata.PackageNotFoundError: result[package] = None
    return {'python':sys.version.split()[0], 'dependencies':result,'network_mode':'public HTTP; no ambient proxy/netrc credentials', 'reasoning':'host model required', 'paid_apis':'not implemented/enabled'}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest='command',required=True)
    commands.add_parser('doctor')
    init = commands.add_parser('init'); init.add_argument('--task',required=True); init.add_argument('--brief',required=True); init.add_argument('--modes',required=True)
    init.add_argument('--depth',choices=['standard','deep'])
    search=commands.add_parser('discover'); search.add_argument('--task',required=True); search.add_argument('--query',required=True); search.add_argument('--limit',type=int,default=8); search.add_argument('--backend',default='auto')
    search.add_argument('--max-age',type=int,default=3600); search.add_argument('--refresh',action='store_true')
    fetch=commands.add_parser('capture'); fetch.add_argument('--task',required=True)
    group=fetch.add_mutually_exclusive_group(required=True); group.add_argument('--url'); group.add_argument('--file')
    fetch.add_argument('--category',default='web'); fetch.add_argument('--selector'); fetch.add_argument('--max-images',type=int,default=30); fetch.add_argument('--timeout',type=int,default=15); fetch.add_argument('--origin-id')
    fetch.add_argument('--source-url');fetch.add_argument('--scope',choices=['full','excerpt'],default='full');fetch.add_argument('--method')
    fetch.add_argument('--max-age',type=int,default=86400);fetch.add_argument('--refresh',action='store_true')
    ev=commands.add_parser('evidence'); ev.add_argument('--task',required=True); ev.add_argument('--source',required=True); ev.add_argument('--block',required=True); ev.add_argument('--quote',required=True); ev.add_argument('--role',default='observation'); ev.add_argument('--context')
    for name in ['validate','finalize']:
        command=commands.add_parser(name);command.add_argument('--task',required=True)
    args=parser.parse_args()
    try:
        if args.command=='doctor': result=doctor()
        elif args.command=='init': result=init_task(args.task,read_json(args.brief),args.modes.split(','),args.depth)
        elif args.command=='discover': result=discover(args.task,args.query,args.limit,args.backend,args.max_age,args.refresh)
        elif args.command=='capture': result=capture(args.task,args.url,args.file,args.category,args.selector,args.max_images,args.timeout,args.origin_id,args.source_url,args.scope,args.method,args.max_age,args.refresh)
        elif args.command=='evidence': result=add_evidence(args.task,args.source,args.block,args.quote,args.role,args.context)
        elif args.command=='validate': result=validate(args.task)
        else: result=finalize(args.task)
        print(json.dumps(result,ensure_ascii=False,indent=2))
        if args.command=='validate' and not result['ok']: return 2
    except Exception as exc:
        print(json.dumps({'status':'failed','error':type(exc).__name__,'message':str(exc)},ensure_ascii=False))
        return 2
    return 0


if __name__=='__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    raise SystemExit(main())
